import io

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from judge.models import ContestParticipation, Organization
from judge.models.tests.util import create_contest, create_contest_problem, create_problem, create_user
from judge.views.generate_users import parse_participants, split_name


class ContestToolsTest(TestCase):
    fixtures = ['language_small.json']

    def setUp(self):
        self.admin = create_user(username='tools_admin', is_superuser=True, is_staff=True)
        self.admin.set_password('tools-pass-1')
        self.admin.save()
        self.client.login(username='tools_admin', password='tools-pass-1')
        now = timezone.now()
        self.contest = create_contest(key='tools', name='Олимпиада', start_time=now - timezone.timedelta(hours=1),
                                      end_time=now + timezone.timedelta(hours=1), is_visible=True,
                                      format_name='ioi16')
        self.cp = create_contest_problem(contest=self.contest, problem=create_problem(code='toolsp', name='НОД'))

    def test_extend(self):
        end = self.contest.end_time
        self.client.post(reverse('contest_extend', args=('tools',)), {'minutes': '15'})
        self.contest.refresh_from_db()
        self.assertEqual(self.contest.end_time - end, timezone.timedelta(minutes=15))
        self.client.post(reverse('contest_extend', args=('tools',)), {'minutes': '999'})
        self.contest.refresh_from_db()
        self.assertEqual(self.contest.end_time - end, timezone.timedelta(minutes=15))

    def test_extend_needs_editor(self):
        student = create_user(username='tools_student')
        student.set_password('tools-pass-1')
        student.save()
        self.client.login(username='tools_student', password='tools-pass-1')
        response = self.client.post(reverse('contest_extend', args=('tools',)), {'minutes': '15'})
        self.assertEqual(response.status_code, 404)

    def test_generate_with_names_universities_and_contest(self):
        response = self.client.post(reverse('admin_generate_users'), {
            'participants': 'Иванов Иван Иванович; ОшГУ\nАбдиллаев Марлен\tОшМПУ\n\nПетрова Анна, ОшГУ\n',
            'prefix': 'tst', 'password_length': 8, 'start_number': 1, 'contest': self.contest.id,
        })
        self.assertEqual(response.status_code, 200)
        users = User.objects.filter(username__startswith='tst').order_by('username')
        self.assertEqual([u.username for u in users], ['tst001', 'tst002', 'tst003'])
        self.assertEqual((users[0].last_name, users[0].first_name), ('Иванов', 'Иван Иванович'))
        self.assertEqual(users[1].profile.organizations.first().name, 'ОшМПУ')
        self.assertEqual(Organization.objects.filter(name='ОшГУ').count(), 1)
        self.assertEqual(Organization.objects.get(name='ОшГУ').slug, 'oshgu')
        self.contest.refresh_from_db()
        self.assertTrue(self.contest.is_private)
        self.assertEqual(self.contest.private_contestants.count(), 3)

    def test_xlsx(self):
        from openpyxl import load_workbook
        student = create_user(username='xl_student', first_name='Айбек', last_name='Токтогулов')
        org = Organization.objects.create(name='ОшГУ', slug='oshsu', short_name='ОшГУ', about='')
        student.profile.organizations.add(org)
        ContestParticipation.objects.create(contest=self.contest, user=student.profile, score=75,
                                            format_data={str(self.cp.id): {'points': 75, 'time': 0}})
        response = self.client.get(reverse('contest_ranking_xlsx', args=('tools',)))
        self.assertEqual(response.status_code, 200)
        ws = load_workbook(io.BytesIO(response.content)).active
        rows = list(ws.iter_rows(values_only=True))
        header = rows[3]
        self.assertIn('НОД', header)
        self.assertEqual(rows[4][:5], (1, 'xl_student', 'Токтогулов Айбек', 'ОшГУ', 75))

    def test_parse_participants(self):
        self.assertEqual(parse_participants('# комментарий\nА Б; ВУЗ\nВ Г\n'), [('А Б', 'ВУЗ'), ('В Г', '')])
        self.assertEqual(split_name('Иванов'), ('', 'Иванов'))
