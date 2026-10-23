"""Tools for contest organizers: results in Excel, extending the contest."""
import io
import re

from django.contrib import messages
from django.http import Http404, HttpResponse, HttpResponseRedirect
from django.shortcuts import get_object_or_404
from django.urls import reverse
from django.utils import timezone
from django.utils.translation import gettext as _
from django.views.decorators.http import require_POST

from judge.models import Contest, ContestParticipation


def ranking_rows(contest):
    problems = list(contest.contest_problems.select_related('problem').order_by('order'))
    participations = (contest.users.filter(virtual=ContestParticipation.LIVE)
                      .select_related('user__user').prefetch_related('user__organizations')
                      .order_by('is_disqualified', '-score', 'cumtime', 'tiebreaker'))
    rows = []
    place = 0
    previous = None
    for index, p in enumerate(participations, 1):
        key = (p.score, p.cumtime, p.tiebreaker)
        if p.is_disqualified:
            shown_place = '—'
        else:
            if key != previous:
                place = index
                previous = key
            shown_place = place
        user = p.user.user
        org = p.user.organization
        data = p.format_data or {}
        cells = []
        for cp in problems:
            cell = data.get(str(cp.id))
            cells.append(cell.get('points') if cell else None)
        rows.append({
            'place': shown_place,
            'username': user.username,
            'name': ' '.join(part for part in (user.last_name, user.first_name) if part),
            'organization': org.name if org else '',
            'problems': cells,
            'score': None if p.is_disqualified else p.score,
            'time': round(p.cumtime / 60) if p.cumtime else 0,
            'status': _('disqualified') if p.is_disqualified else '',
        })
    return problems, rows


def contest_ranking_xlsx(request, contest):
    contest = get_object_or_404(Contest, key=contest)
    if not request.user.is_authenticated or not contest.is_editable_by(request.user):
        raise Http404()

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    problems, rows = ranking_rows(contest)
    wb = Workbook()
    ws = wb.active
    ws.title = _('Results')[:31]

    ws.append([contest.name])
    ws['A1'].font = Font(bold=True, size=14)
    when = timezone.localtime(contest.start_time)
    ws.append(['%s %s' % (_('Date:'), when.strftime('%d.%m.%Y %H:%M'))])
    ws.append([])

    header = [_('Place'), _('Login'), _('Full name'), _('University / school')]
    header += [cp.problem.name for cp in problems]
    header += [_('Total'), _('Time, min'), _('Note')]
    ws.append(header)
    head_row = ws.max_row
    for cell in ws[head_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='DCE6F1')
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')

    for row in rows:
        ws.append([row['place'], row['username'], row['name'], row['organization']] + row['problems'] +
                  [row['score'], row['time'], row['status']])

    widths = [7, 18, 32, 26] + [14] * len(problems) + [10, 10, 16]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = ws.cell(row=head_row + 1, column=3)

    out = io.BytesIO()
    wb.save(out)
    name = re.sub(r'[^A-Za-z0-9_-]+', '-', contest.key) or 'contest'
    response = HttpResponse(out.getvalue(),
                            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="results-%s.xlsx"' % name
    return response


@require_POST
def contest_extend(request, contest):
    contest = get_object_or_404(Contest, key=contest)
    if not request.user.is_authenticated or not contest.is_editable_by(request.user):
        raise Http404()
    try:
        minutes = int(request.POST.get('minutes', 0))
    except ValueError:
        minutes = 0
    if not -180 <= minutes <= 180 or minutes == 0:
        messages.error(request, _('Enter a number of minutes from -180 to 180.'))
    else:
        contest.end_time += timezone.timedelta(minutes=minutes)
        if contest.end_time <= contest.start_time:
            messages.error(request, _('The contest cannot end before it starts.'))
        else:
            contest.save(update_fields=['end_time'])
            messages.success(request, _('The contest now ends at %s.') %
                             timezone.localtime(contest.end_time).strftime('%d.%m.%Y %H:%M'))
    return HttpResponseRedirect(reverse('contest_view', args=(contest.key,)))
