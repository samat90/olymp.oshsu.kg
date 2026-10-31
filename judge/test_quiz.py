import itertools
import json
import os
import shutil
import tempfile
from io import BytesIO
from unittest import skipUnless
from unittest.mock import patch

from django.core import signing
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.template import TemplateDoesNotExist
from django.template.loader import get_template
from django.test import Client, TestCase, override_settings
from django.urls import reverse

from judge.models import GameSession, MiscConfig, Quiz, QuizAnswer, QuizOption, QuizPlayer, QuizQuestion
from judge.models.tests.util import create_user
from judge.utils import quiz_game as engine
from judge.utils import quiz_scoring as scoring
from judge.views.quiz import COOKIE_SALT, player_cookie, quiz_records, visible_quizzes

Q = QuizQuestion
KEYS = itertools.count(1)
LEAK_KEYS = {'is_correct', 'correct', 'answers', 'word', 'answer', 'x', 'y', 'r'}


def template_exists(*names):
    for name in names:
        try:
            get_template(name)
        except TemplateDoesNotExist:
            return False
    return True


def needs_templates(*names):
    """Page tests turn on by themselves once the templates exist."""
    return skipUnless(template_exists(*names), 'templates not written yet: %s' % ', '.join(names))


def make_quiz(author, title='Викторина', access=Quiz.PUBLIC, **kwargs):
    return Quiz.objects.create(author=author.profile, title=title, access=access, key='t%07d' % next(KEYS), **kwargs)


def add_question(quiz, type, options=(), **kwargs):
    question = Q.objects.create(quiz=quiz, type=type, order=quiz.questions.count(), text='Вопрос %s' % type,
                                **kwargs)
    for n, (text, ok) in enumerate(options):
        QuizOption.objects.create(question=question, order=n, text=text, is_correct=ok)
    return question


def option_id(question, text):
    return question.options.get(text=text).id


def walk_keys(value):
    if isinstance(value, dict):
        for key, item in value.items():
            yield key
            yield from walk_keys(item)
    elif isinstance(value, list):
        for item in value:
            yield from walk_keys(item)


# ---------------------------------------------------------------- checking and points

class QuizScoringTest(TestCase):
    def setUp(self):
        cache.clear()
        self.author = create_user(username='scorer')
        self.quiz = make_quiz(self.author)

    def key(self, type, options=(), **kwargs):
        question = add_question(self.quiz, type, options, **kwargs)
        return question, scoring.build_key(question)

    def result(self, key, value):
        cleaned = scoring.clean_value(key, value)
        self.assertIsNotNone(cleaned, 'value %r was rejected' % (value,))
        return scoring.check(key, cleaned)

    def test_single_and_truefalse(self):
        q, key = self.key(Q.SINGLE, [('a', True), ('b', False), ('c', False)])
        self.assertEqual(self.result(key, option_id(q, 'a')), (1.0, True))
        self.assertEqual(self.result(key, str(option_id(q, 'b'))), (0.0, False))
        self.assertIsNone(scoring.clean_value(key, 999999))
        self.assertIsNone(scoring.clean_value(key, 'abc'))
        self.assertIsNone(scoring.clean_value(key, None))
        q, key = self.key(Q.TRUEFALSE, [('Верно', False), ('Неверно', True)])
        self.assertEqual(self.result(key, option_id(q, 'Неверно')), (1.0, True))
        self.assertEqual(self.result(key, option_id(q, 'Верно')), (0.0, False))

    def test_multi_partial(self):
        q, key = self.key(Q.MULTI, [('a', True), ('b', True), ('c', False), ('d', False)])
        a, b, c, d = (option_id(q, x) for x in 'abcd')
        self.assertEqual(self.result(key, [a, b]), (1.0, True))
        self.assertEqual(self.result(key, [b, a, a]), (1.0, True))
        self.assertEqual(self.result(key, [a]), (0.5, False))
        self.assertEqual(self.result(key, [a, b, c]), (0.5, False))   # right minus extra
        self.assertEqual(self.result(key, [a, c]), (0.0, False))
        self.assertEqual(self.result(key, [c, d]), (0.0, False))      # never below zero
        self.assertEqual(self.result(key, []), (0.0, False))
        self.assertIsNone(scoring.clean_value(key, [a, 999999]))
        self.assertIsNone(scoring.clean_value(key, 'ab'))

    def test_word_ignores_case_and_yo(self):
        _q, key = self.key(Q.WORD, data={'word': 'Ёлка'})
        self.assertEqual(self.result(key, 'елка'), (1.0, True))
        self.assertEqual(self.result(key, '  ЁЛКА '), (1.0, True))
        self.assertEqual(self.result(key, 'ЛЕКА'), (0.0, False))
        self.assertIsNone(scoring.clean_value(key, '   '))

    def test_order_is_exact_only(self):
        q, key = self.key(Q.ORDER, [('один', False), ('два', False), ('три', False), ('четыре', False)])
        ids = [o.id for o in q.options.all()]
        self.assertEqual(self.result(key, ids), (1.0, True))
        self.assertEqual(self.result(key, [ids[1], ids[0]] + ids[2:]), (0.0, False))
        self.assertEqual(self.result(key, ids[::-1]), (0.0, False))
        self.assertIsNone(scoring.clean_value(key, ids[:3]))                 # an id is missing
        self.assertIsNone(scoring.clean_value(key, ids[:3] + [ids[0]]))      # repeated id
        self.assertIsNone(scoring.clean_value(key, ids + [999999]))

    def test_slider_closeness(self):
        _q, key = self.key(Q.SLIDER, data={'min': 0, 'max': 100, 'step': 1, 'answer': 42, 'tolerance': 2})
        self.assertEqual(self.result(key, 42), (1.0, True))
        self.assertEqual(self.result(key, 44), (1.0, True))
        fraction, ok = self.result(key, '39,5')
        self.assertAlmostEqual(fraction, 0.975)
        self.assertFalse(ok)
        fraction, ok = self.result(key, 45)
        self.assertFalse(ok)
        self.assertAlmostEqual(fraction, 0.95)
        self.assertEqual(self.result(key, 80), (0.0, False))
        self.assertEqual(scoring.clean_value(key, 150), 100)    # clamped to the scale
        self.assertIsNone(scoring.clean_value(key, float('nan')))
        self.assertIsNone(scoring.clean_value(key, float('inf')))
        self.assertIsNone(scoring.clean_value(key, 'nan'))
        self.assertIsNone(scoring.clean_value(key, 'много'))
        self.assertIsNone(scoring.clean_value(key, [42]))

    def test_text_answers_and_numbers(self):
        _q, key = self.key(Q.TEXT, data={'answers': ['Бишкек', 'г. Бишкек'], 'tolerance': 0})
        self.assertEqual(self.result(key, ' бишкек '), (1.0, True))
        self.assertEqual(self.result(key, 'Г.  БИШКЕК'), (1.0, True))
        self.assertEqual(self.result(key, 'Ош'), (0.0, False))
        self.assertIsNone(scoring.clean_value(key, ''))
        self.assertIsNone(scoring.clean_value(key, '   '))
        self.assertIsNone(scoring.clean_value(key, ['Бишкек']))
        _q, key = self.key(Q.TEXT, data={'answers': ['3.14'], 'tolerance': 0.01})
        self.assertEqual(self.result(key, '3,145'), (1.0, True))
        self.assertEqual(self.result(key, '3.13'), (1.0, True))
        self.assertEqual(self.result(key, '3.2'), (0.0, False))
        _q, key = self.key(Q.TEXT, data={'answers': ['10'], 'tolerance': 0})
        self.assertEqual(self.result(key, '10,0'), (1.0, True))
        self.assertEqual(self.result(key, '10.5'), (0.0, False))

    def test_code_with_options_and_typed(self):
        q, key = self.key(Q.CODE, [('5', True), ('23', False)], code='print(2 + 3)', code_language='python',
                          data={'input': False})
        self.assertEqual(self.result(key, option_id(q, '5')), (1.0, True))
        self.assertEqual(self.result(key, option_id(q, '23')), (0.0, False))
        self.assertIsNone(scoring.clean_value(key, 'x5x'))
        _q, key = self.key(Q.CODE, code='print(2 + 3)', data={'input': True, 'answers': ['5'], 'tolerance': 0})
        self.assertEqual(self.result(key, ' 5 '), (1.0, True))
        self.assertEqual(self.result(key, '5,0'), (1.0, True))
        self.assertEqual(self.result(key, '23'), (0.0, False))

    def test_pin_radius(self):
        _q, key = self.key(Q.PIN, data={'x': 30, 'y': 60, 'r': 8})
        self.assertEqual(self.result(key, {'x': 35, 'y': 65}), (1.0, True))
        self.assertEqual(self.result(key, {'x': '30', 'y': '52'}), (1.0, True))
        self.assertEqual(self.result(key, {'x': 40, 'y': 60}), (0.0, False))
        self.assertEqual(scoring.clean_value(key, {'x': -5, 'y': 120}), {'x': 0, 'y': 100})
        self.assertIsNone(scoring.clean_value(key, {'x': 'a', 'y': 1}))
        self.assertIsNone(scoring.clean_value(key, {'x': 1}))
        self.assertIsNone(scoring.clean_value(key, '30,60'))
        self.assertIsNone(scoring.clean_value(key, {'x': float('nan'), 'y': 1}))

    def test_unscored_types(self):
        q, key = self.key(Q.POLL, [('x', False), ('y', False)])
        self.assertEqual(self.result(key, option_id(q, 'x')), (0.0, False))
        self.assertFalse(scoring.scored(key))
        self.assertEqual(scoring.points(key, 1.0, True, 0, 3), 0)
        _q, key = self.key(Q.CLOUD)
        self.assertEqual(scoring.clean_value(key, 'а' * 50), 'а' * 30)
        self.assertEqual(self.result(key, 'слово'), (0.0, False))
        self.assertFalse(scoring.scored(key))
        _q, key = self.key(Q.OPEN)
        self.assertEqual(len(scoring.clean_value(key, 'б' * 500)), 300)
        self.assertEqual(self.result(key, 'мысль'), (0.0, False))
        self.assertFalse(scoring.scored(key))
        self.assertFalse(Q(type=Q.OPEN, weight=1).scored)
        self.assertTrue(Q(type=Q.SINGLE, weight=2).scored)

    def test_points_formula(self):
        _q, key = self.key(Q.SINGLE, [('a', True), ('b', False)], time_limit=20)
        self.assertEqual(scoring.points(key, 1.0, True, 0, 0), 1000)
        self.assertEqual(scoring.points(key, 1.0, True, 10, 0), 750)
        self.assertEqual(scoring.points(key, 1.0, True, 20, 0), 500)
        self.assertEqual(scoring.points(key, 1.0, True, 99, 0), 500)     # never below half
        self.assertEqual(scoring.points(key, 1.0, True, -5, 0), 1000)
        self.assertEqual(scoring.points(key, 0.0, False, 0, 3), 0)
        self.assertEqual(scoring.points(key, 0.5, False, 0, 3), 500)     # partial: no streak bonus
        _q, double = self.key(Q.SINGLE, [('a', True)], time_limit=20, weight=2)
        self.assertEqual(scoring.points(double, 1.0, True, 10, 0), 1500)
        _q, zero = self.key(Q.SINGLE, [('a', True)], time_limit=20, weight=0)
        self.assertEqual(scoring.points(zero, 1.0, True, 0, 4), 0)

    def test_streak_bonus(self):
        _q, key = self.key(Q.SINGLE, [('a', True)], time_limit=20)
        self.assertEqual(scoring.points(key, 1.0, True, 0, 1), 1100)
        self.assertEqual(scoring.points(key, 1.0, True, 0, 3), 1300)
        self.assertEqual(scoring.points(key, 1.0, True, 0, 5), 1500)
        self.assertEqual(scoring.points(key, 1.0, True, 0, 12), 1500)    # capped at +500
        self.assertEqual(scoring.points(key, 0.0, False, 0, 12), 0)


# ---------------------------------------------------------------- games over HTTP

class QuizGameTestCase(TestCase):
    """Fixed server clock (judge.utils.quiz_game.now) and no countdown."""

    def setUp(self):
        cache.clear()
        self.t = 1800000000.0
        for p in (patch.object(engine, 'now', lambda: self.t), patch.object(engine, 'COUNTDOWN', 0)):
            p.start()
            self.addCleanup(p.stop)
        self.author = create_user(username='quizhost')
        self.other = create_user(username='stranger')
        self.admin = create_user(username='quizadmin', is_superuser=True, is_staff=True)
        self.student = create_user(username='pupil')
        self.host = self.login(self.author)
        self.quiz = make_quiz(self.author, title='Столицы')
        self.q1 = add_question(self.quiz, Q.SINGLE, [('Два', True), ('Три', False), ('Пять', False)])
        self.q2 = add_question(self.quiz, Q.SINGLE, [('Да', True), ('Нет', False)])

    # helpers

    def login(self, user, **kwargs):
        client = Client(**kwargs)
        client.force_login(user)
        return client

    def launch(self, quiz=None, mode='screen', client=None, **data):
        quiz = quiz or self.quiz
        response = (client or self.host).post(reverse('quiz_launch', args=[quiz.key]), dict(data, mode=mode))
        self.assertEqual(response.status_code, 302)
        game = GameSession.objects.latest('id')
        self.assertEqual(response['Location'], reverse('quiz_host', args=[game.pin]))
        return game.pin

    def game(self, pin):
        return GameSession.objects.filter(pin=pin).latest('id')

    def join(self, pin, nickname, client=None, status=302):
        client = client or Client()
        response = client.post(reverse('quiz_join'), {'pin': pin, 'nickname': nickname})
        self.assertEqual(response.status_code, status, response.content[:300])
        if status == 302:
            self.assertEqual(response['Location'], reverse('quiz_play', args=[pin]))
        return client

    def player(self, pin, nickname):
        return QuizPlayer.objects.get(game=self.game(pin), nickname=nickname)

    def state(self, client, pin, v=None):
        url = reverse('quiz_state', args=[pin]) + ('?v=%s' % v if v is not None else '')
        response = client.get(url)
        self.assertEqual(response.status_code, 200)
        return response.json()

    def question(self, client, pin):
        response = client.get(reverse('quiz_question', args=[pin]))
        self.assertEqual(response.status_code, 200)
        return response.json()

    def answer(self, client, pin, i, value, **extra):
        body = dict(extra, i=i, value=value)
        return client.post(reverse('quiz_answer', args=[pin]), json.dumps(body), content_type='application/json')

    def result(self, client, pin):
        return client.get(reverse('quiz_result', args=[pin]))

    def action(self, pin, action, client=None, status=200, **data):
        response = (client or self.host).post(reverse('quiz_host_action', args=[pin]), dict(data, action=action))
        self.assertEqual(response.status_code, status, response.content[:300])
        return response.json()

    def phase(self, pin):
        return self.state(Client(), pin)['ph']

    def game_with_players(self, *nicknames, mode='screen', quiz=None):
        pin = self.launch(quiz=quiz, mode=mode)
        return pin, [self.join(pin, nick) for nick in nicknames]


class QuizTimingTest(QuizGameTestCase):
    def test_time_counts_from_when_the_question_was_fetched(self):
        pin, (a, b, c) = self.game_with_players('Аня', 'Бек', 'Чынара')
        self.action(pin, 'start')
        start = self.t
        self.t = start + 4
        self.question(a, pin)                       # A sees the question 4 s late
        self.question(c, pin)                       # C is here too, so the question stays open
        self.t = start + 6
        right = option_id(self.q1, 'Два')
        # whatever the client says about time is ignored
        self.assertEqual(self.answer(a, pin, 0, right, time=0, elapsed=0, ms=0).status_code, 200)
        self.assertEqual(self.answer(b, pin, 0, right, time=0).status_code, 200)   # B never fetched it
        a_answer = QuizAnswer.objects.get(player=self.player(pin, 'Аня'))
        b_answer = QuizAnswer.objects.get(player=self.player(pin, 'Бек'))
        self.assertEqual(a_answer.time_ms, 2000)
        self.assertEqual(a_answer.points, 950)
        self.assertEqual(b_answer.time_ms, 6000)
        self.assertEqual(b_answer.points, 850)

    def test_fetching_during_the_countdown_counts_from_the_start(self):
        with patch.object(engine, 'COUNTDOWN', 3):
            pin, (a, b) = self.game_with_players('Аня', 'Бек')
            self.action(pin, 'start')
            opened = self.t
            self.question(a, pin)                   # during the countdown
            self.t = opened + 3 + 1
            self.assertEqual(self.answer(a, pin, 0, option_id(self.q1, 'Два')).status_code, 200)
        answer = QuizAnswer.objects.get()
        self.assertEqual(answer.time_ms, 1000)
        self.assertEqual(answer.points, 975)

    def test_answer_during_the_countdown_is_rejected(self):
        # The question is already served during the countdown; an answer sent before `start`
        # must not get full points.
        with patch.object(engine, 'COUNTDOWN', 3):
            pin, (a, b) = self.game_with_players('Аня', 'Бек')
            self.action(pin, 'start')
            self.question(a, pin)
            self.t += 1
            response = self.answer(a, pin, 0, option_id(self.q1, 'Два'))
        self.assertEqual(response.status_code, 409)
        self.assertFalse(QuizAnswer.objects.exists())

    def test_answer_after_deadline_is_rejected(self):
        pin, (a, b) = self.game_with_players('Аня', 'Бек')
        self.action(pin, 'start')
        end = self.t + self.q1.time_limit
        self.t = end + engine.GRACE - 0.1
        self.assertEqual(self.answer(a, pin, 0, option_id(self.q1, 'Два')).status_code, 200)
        self.t = end + engine.GRACE + 0.1
        self.assertEqual(self.answer(b, pin, 0, option_id(self.q1, 'Два')).status_code, 409)
        self.assertEqual(QuizAnswer.objects.count(), 1)
        self.assertEqual(QuizAnswer.objects.get().points, 500)


class QuizPermissionRulesTest(QuizGameTestCase):
    """Anyone signed in creates and hosts any catalog quiz; only the author edits or deletes it."""

    def test_new_quiz_is_in_the_catalog_by_default(self):
        response = self.login(self.other).get(reverse('quiz_new'))
        quiz = Quiz.objects.get(author=self.other.profile)
        self.assertRedirects(response, reverse('quiz_edit', args=[quiz.key]), fetch_redirect_response=False)
        self.assertEqual(quiz.access, Quiz.PUBLIC)
        add_question(quiz, Q.SINGLE, [('Да', True), ('Нет', False)])
        self.assertIn(quiz, visible_quizzes())

    def test_anyone_signed_in_hosts_someone_elses_quiz_in_every_mode(self):
        stranger = self.login(self.other)
        for mode in ('screen', 'auto'):
            pin = self.launch(mode=mode, client=stranger)
            self.assertEqual(self.game(pin).host, self.other.profile)
            self.assertEqual(stranger.get(reverse('quiz_host', args=[pin])).status_code, 200)
        solo = stranger.post(reverse('quiz_solo', args=[self.quiz.key]))
        self.assertEqual(solo.status_code, 302)
        self.assertEqual(Client().post(reverse('quiz_launch', args=[self.quiz.key]), {'mode': 'screen'})
                         .status_code, 302)                     # to the login page, no game
        self.assertFalse(GameSession.objects.filter(host=None).exclude(mode='solo').exists())

    def test_only_the_author_edits_and_deletes(self):
        q = self.q1
        api = 'quiz_api_question_save', 'quiz_api_question_delete', 'quiz_api_question_image'
        staff = create_user(username='plainstaff', is_staff=True)       # staff without the quiz permission
        for user in (self.other, staff):
            client = self.login(user)
            self.assertIn(client.get(reverse('quiz_edit', args=[self.quiz.key])).status_code, (403, 404))
            self.assertIn(client.post(reverse('quiz_delete', args=[self.quiz.key])).status_code, (403, 404))
            for name in api:
                self.assertEqual(client.post(reverse(name, args=[self.quiz.key, q.id]), '{}',
                                             content_type='application/json').status_code, 404)
            for name in ('quiz_api_save', 'quiz_api_reorder', 'quiz_api_import', 'quiz_api_cover',
                         'quiz_api_question_add'):
                self.assertEqual(client.post(reverse(name, args=[self.quiz.key]), '{}',
                                             content_type='application/json').status_code, 404)
            page = client.get(reverse('quiz_detail', args=[self.quiz.key])).content.decode()
            self.assertNotIn(reverse('quiz_edit', args=[self.quiz.key]), page)
        self.assertTrue(Quiz.objects.filter(id=self.quiz.id, is_deleted=False).exists())
        self.assertEqual(self.quiz.questions.count(), 2)
        page = self.host.get(reverse('quiz_detail', args=[self.quiz.key])).content.decode()
        self.assertIn(reverse('quiz_edit', args=[self.quiz.key]), page)
        self.assertEqual(self.host.get(reverse('quiz_edit', args=[self.quiz.key])).status_code, 200)

    def test_site_admins_manage_every_quiz(self):
        from django.contrib.auth.models import Permission
        moderator = create_user(username='moderator', is_staff=True)
        moderator.user_permissions.add(Permission.objects.get(codename='change_quiz', content_type__app_label='judge'))
        for user in (self.admin, moderator):
            client = self.login(user)
            self.assertEqual(client.get(reverse('quiz_edit', args=[self.quiz.key])).status_code, 200)
            page = client.get(reverse('quiz_detail', args=[self.quiz.key])).content.decode()
            self.assertIn(reverse('quiz_edit', args=[self.quiz.key]), page)
            response = client.post(reverse('quiz_api_save', args=[self.quiz.key]),
                                   json.dumps({'title': 'Правка админа %s' % user.username}),
                                   content_type='application/json')
            self.assertEqual(response.status_code, 200)
            added = client.post(reverse('quiz_api_question_add', args=[self.quiz.key]), json.dumps({'type': 'poll'}),
                                content_type='application/json')
            self.assertEqual(added.status_code, 200)
            qid = added.json()['question']['id']
            self.assertEqual(client.post(reverse('quiz_api_question_delete', args=[self.quiz.key, qid]), '{}',
                                         content_type='application/json').status_code, 200)
        self.quiz.refresh_from_db()
        self.assertEqual(self.quiz.title, 'Правка админа moderator')
        pin = self.launch()
        game = self.game(pin)
        self.assertEqual(self.login(moderator).get(reverse('quiz_game_report', args=[game.id])).status_code, 200)
        self.assertEqual(self.login(self.admin).post(reverse('quiz_delete', args=[self.quiz.key])).status_code, 302)
        self.quiz.refresh_from_db()
        self.assertTrue(self.quiz.is_deleted)                         # it was played: reports stay

    def test_admin_moderates_by_hiding(self):
        self.quiz.is_hidden = True
        self.quiz.save()
        self.assertEqual(self.login(self.other).get(reverse('quiz_detail', args=[self.quiz.key])).status_code, 404)
        self.assertEqual(self.login(self.admin).get(reverse('quiz_detail', args=[self.quiz.key])).status_code, 200)
        self.assertNotIn(self.quiz, visible_quizzes())

    def test_deleting_a_played_quiz_keeps_other_hosts_reports(self):
        stranger = self.login(self.other)
        pin = self.launch(client=stranger)
        a = self.join(pin, 'Аня')
        self.action(pin, 'start', client=stranger)
        self.answer(a, pin, 0, option_id(self.q1, 'Два'))
        self.action(pin, 'end', client=stranger)
        game = self.game(pin)
        response = self.host.post(reverse('quiz_delete', args=[self.quiz.key]))
        self.assertEqual(response.status_code, 302)
        self.quiz.refresh_from_db()
        self.assertTrue(self.quiz.is_deleted)
        self.assertEqual(self.host.get(reverse('quiz_detail', args=[self.quiz.key])).status_code, 404)
        self.assertEqual(self.host.get(reverse('quiz_edit', args=[self.quiz.key])).status_code, 404)
        self.assertNotIn(self.quiz, visible_quizzes())
        report = stranger.get(reverse('quiz_game_report', args=[game.id]))
        self.assertEqual(report.status_code, 200)
        self.assertIn('Аня', report.content.decode())
        self.assertEqual(stranger.get(reverse('quiz_game_report_xlsx', args=[game.id])).status_code, 200)

    def test_deleting_an_unplayed_quiz_removes_it(self):
        quiz = make_quiz(self.author, title='Черновик')
        self.assertEqual(self.host.post(reverse('quiz_delete', args=[quiz.key])).status_code, 302)
        self.assertFalse(Quiz.objects.filter(id=quiz.id).exists())


class QuizEarlyCloseTest(QuizGameTestCase):
    def test_long_question_closes_when_everyone_answered(self):
        self.q1.time_limit = 50
        self.q1.save()
        pin, (a, b) = self.game_with_players('Аня', 'Бек')
        self.action(pin, 'start')
        self.t += 8
        self.answer(a, pin, 0, option_id(self.q1, 'Два'))
        self.assertEqual(self.phase(pin), 'question')
        self.answer(b, pin, 0, option_id(self.q1, 'Три'))
        self.assertEqual(self.phase(pin), 'reveal')                 # at 8 s of 50, no grace wait

    def test_a_closed_tab_does_not_hold_the_question(self):
        pin, (a, b, gone) = self.game_with_players('Аня', 'Бек', 'Ушёл')
        self.action(pin, 'start')
        self.question(a, pin)                                       # these two received the question,
        self.question(b, pin)                                       # the third player never polled again
        self.t += engine.SEEN_WINDOW + 1
        self.answer(a, pin, 0, option_id(self.q1, 'Два'))
        self.assertEqual(self.phase(pin), 'question')
        self.answer(b, pin, 0, option_id(self.q1, 'Два'))
        self.assertEqual(self.phase(pin), 'reveal')

    def test_auto_mode_closes_early_too(self):
        pin, (a,) = self.game_with_players('Аня', mode='auto')
        self.action(pin, 'start')
        self.answer(a, pin, 0, option_id(self.q1, 'Два'))
        self.assertEqual(self.phase(pin), 'reveal')


class QuizDeletedQuestionTest(QuizGameTestCase):
    def test_report_keeps_answers_of_a_deleted_question(self):
        pin, (a, b) = self.game_with_players('Аня', 'Бек')
        self.action(pin, 'start')
        self.answer(a, pin, 0, option_id(self.q1, 'Два'))
        self.answer(b, pin, 0, option_id(self.q1, 'Три'))
        self.action(pin, 'end')
        game = self.game(pin)
        self.q1.delete()
        self.assertEqual(QuizAnswer.objects.filter(player__game=game).count(), 2)
        page = self.host.get(reverse('quiz_game_report', args=[game.id]))
        self.assertEqual(page.status_code, 200)
        self.assertIn('удалён', page.content.decode())
        from judge.views.quiz import report_rows
        questions, players, answers, stats = report_rows(game)
        self.assertEqual(stats[0]['answered'], 2)
        self.assertEqual(stats[0]['correct'], 1)
        self.assertEqual(self.host.get(reverse('quiz_game_report_xlsx', args=[game.id])).status_code, 200)


class QuizLeakTest(QuizGameTestCase):
    def all_types_quiz(self):
        quiz = make_quiz(self.author, title='Все типы')
        add_question(quiz, Q.SINGLE, [('Два', True), ('Три', False), ('Пять', False)])
        add_question(quiz, Q.TRUEFALSE, [('Верно', True), ('Неверно', False)])
        add_question(quiz, Q.MULTI, [('a', True), ('b', True), ('c', False), ('d', False)])
        add_question(quiz, Q.WORD, data={'word': 'АЛГОРИТМ'})
        add_question(quiz, Q.ORDER, [('один', False), ('два', False), ('три', False), ('четыре', False)])
        add_question(quiz, Q.CODE, [('5', True), ('23', False)], code='print(2 + 3)', code_language='python',
                     data={'input': False})
        add_question(quiz, Q.CODE, code='print(2 + 3)', code_language='python',
                     data={'input': True, 'answers': ['5'], 'tolerance': 0})
        add_question(quiz, Q.TEXT, data={'answers': ['Бишкек', 'г. Бишкек'], 'tolerance': 0})
        add_question(quiz, Q.SLIDER, data={'min': 0, 'max': 100, 'step': 1, 'answer': 42, 'tolerance': 2,
                                           'unit': 'км'})
        add_question(quiz, Q.PIN, data={'x': 30.5, 'y': 60, 'r': 8}, image='quiz/map.jpg')
        add_question(quiz, Q.POLL, [('x', False), ('y', False)])
        add_question(quiz, Q.CLOUD)
        add_question(quiz, Q.OPEN)
        return quiz

    def test_unfinished_questions_block_the_launch(self):
        quiz = make_quiz(self.author, title='Недописанная')
        add_question(quiz, Q.SINGLE, [('Два', False), ('Три', False)])
        add_question(quiz, Q.PIN, data={'x': 10, 'y': 10, 'r': 5})
        response = self.host.post(reverse('quiz_launch', args=[quiz.key]), {'mode': 'screen'})
        self.assertEqual(response.status_code, 400)
        self.assertIn('1, 2', response.content.decode())
        self.assertFalse(GameSession.objects.filter(quiz=quiz).exists())

    def test_question_never_contains_the_answer(self):
        quiz = self.all_types_quiz()
        questions = list(quiz.questions.order_by('order'))
        pin, (a, b) = self.game_with_players('Аня', 'Бек', quiz=quiz)
        self.action(pin, 'start')
        seen_types = set()
        for i, q in enumerate(questions):
            self.assertEqual(self.phase(pin), 'question')
            for client in (a, Client()):
                payload = self.question(client, pin)
                self.assertEqual(payload['i'], i)
                self.assertEqual(payload['type'], q.type)
                leaked = LEAK_KEYS & set(walk_keys(payload))
                self.assertFalse(leaked, '%s question shows %s' % (q.type, leaked))
                text = json.dumps(payload, ensure_ascii=False)
                if q.type == Q.WORD:
                    self.assertEqual(sorted(payload['letters']), sorted('АЛГОРИТМ'))
                    self.assertNotEqual(''.join(payload['letters']), 'АЛГОРИТМ')
                    self.assertNotIn('АЛГОРИТМ', text.upper())
                if q.type == Q.ORDER:
                    self.assertNotEqual([o['id'] for o in payload['options']], [o.id for o in q.options.all()])
                if q.type == Q.TEXT:
                    self.assertNotIn('Бишкек', text)
            seen_types.add(q.type)
            self.assertEqual(self.result(a, pin).status_code, 409)
            self.assertEqual(self.result(Client(), pin).status_code, 409)
            self.action(pin, 'skip')
            reveal = self.result(a, pin).json()['reveal']
            self.assertEqual(reveal['type'], q.type)
            self.action(pin, 'next')
            self.action(pin, 'next')
        self.assertEqual(len(seen_types), 12)
        self.assertEqual(self.phase(pin), 'finished')


class QuizAnswerRulesTest(QuizGameTestCase):
    def test_one_answer_per_question(self):
        pin, (a, b) = self.game_with_players('Аня', 'Бек')
        self.action(pin, 'start')
        self.assertEqual(self.answer(a, pin, 0, option_id(self.q1, 'Три')).status_code, 200)
        response = self.answer(a, pin, 0, option_id(self.q1, 'Два'))
        self.assertEqual(response.status_code, 409)
        self.assertEqual(QuizAnswer.objects.count(), 1)
        self.assertFalse(QuizAnswer.objects.get().is_correct)
        self.assertTrue(self.question(a, pin)['answered'])
        self.assertFalse(self.question(b, pin)['answered'])

    def test_wrong_question_number_and_bad_values(self):
        pin, (a, b) = self.game_with_players('Аня', 'Бек')
        self.action(pin, 'start')
        self.assertEqual(self.answer(a, pin, 1, option_id(self.q1, 'Два')).status_code, 409)
        self.assertEqual(self.answer(a, pin, 0, option_id(self.q2, 'Да')).status_code, 400)
        self.assertEqual(self.answer(a, pin, 0, 'x').status_code, 400)
        url = reverse('quiz_answer', args=[pin])
        self.assertEqual(a.post(url, 'not json', content_type='application/json').status_code, 400)
        self.assertEqual(a.post(url, '[1, 2]', content_type='application/json').status_code, 400)
        self.assertEqual(a.post(url, '{"i": 0}', content_type='application/json').status_code, 400)
        self.assertEqual(a.get(url).status_code, 405)
        self.assertFalse(QuizAnswer.objects.exists())

    def test_answer_needs_csrf_token(self):
        pin = self.launch()
        client = Client(enforce_csrf_checks=True)
        client.get(reverse('quiz_join') + '?pin=' + pin)
        token = client.cookies['csrftoken'].value
        response = client.post(reverse('quiz_join'), {'pin': pin, 'nickname': 'Аня', 'csrfmiddlewaretoken': token})
        self.assertEqual(response.status_code, 302)
        self.join(pin, 'Бек')
        self.action(pin, 'start')
        url = reverse('quiz_answer', args=[pin])
        body = json.dumps({'i': 0, 'value': option_id(self.q1, 'Два')})
        response = client.post(url, body, content_type='application/json')
        self.assertEqual(response.status_code, 403)
        self.assertFalse(QuizAnswer.objects.exists())
        response = client.post(url, body, content_type='application/json', HTTP_X_CSRFTOKEN=token)
        self.assertEqual(response.status_code, 200)

    def test_answer_without_player_cookie(self):
        pin, (a, b) = self.game_with_players('Аня', 'Бек')
        self.action(pin, 'start')
        self.assertEqual(self.answer(Client(), pin, 0, option_id(self.q1, 'Два')).status_code, 403)
        self.assertEqual(self.answer(self.login(self.other), pin, 0, option_id(self.q1, 'Два')).status_code, 403)
        self.assertFalse(QuizAnswer.objects.exists())

    def test_cookie_from_another_game(self):
        pin_a, (a,) = self.game_with_players('Аня')
        pin_b, (b,) = self.game_with_players('Бек')
        self.action(pin_b, 'start')
        player = self.player(pin_a, 'Аня')
        name = player_cookie(pin_b)
        intruder = Client()
        # a correctly signed cookie for game B that names game A
        intruder.cookies[name] = signing.get_cookie_signer(salt=name + COOKIE_SALT).sign(
            '%d:%d' % (self.game(pin_a).id, player.id))
        self.assertEqual(self.answer(intruder, pin_b, 0, option_id(self.q1, 'Два')).status_code, 403)
        self.assertFalse(QuizAnswer.objects.exists())

    def test_tampered_cookie_is_ignored(self):
        # A cookie copied under another game's name has a bad signature: it must count as no cookie, not a 500.
        pin_a, (a,) = self.game_with_players('Аня')
        pin_b, (b,) = self.game_with_players('Бек')
        self.action(pin_b, 'start')
        intruder = Client()
        intruder.cookies[player_cookie(pin_b)] = a.cookies[player_cookie(pin_a)].value
        self.assertEqual(self.answer(intruder, pin_b, 0, option_id(self.q1, 'Два')).status_code, 403)
        intruder.cookies[player_cookie(pin_b)] = 'garbage'
        self.assertEqual(self.state(intruder, pin_b)['ph'], 'question')
        self.assertFalse(QuizAnswer.objects.exists())


class QuizPlayersTest(QuizGameTestCase):
    def test_guest_and_user_join(self):
        pin = self.launch()
        guest = self.join(pin, '  Аня   Ким ')
        user = self.login(self.student)
        response = user.get(reverse('quiz_join') + '?pin=' + pin)     # users join with their username
        self.assertRedirects(response, reverse('quiz_play', args=[pin]), fetch_redirect_response=False)
        players = {p.nickname: p for p in QuizPlayer.objects.filter(game=self.game(pin))}
        self.assertEqual(set(players), {'Аня Ким', 'pupil'})
        self.assertEqual(players['pupil'].profile, self.student.profile)
        self.assertIsNone(players['Аня Ким'].profile)
        self.assertEqual(len(players['Аня Ким'].token), 32)
        self.assertEqual(engine.player_count(self.game(pin).id), 2)
        host_state = self.host.get(reverse('quiz_host_state', args=[pin])).json()
        self.assertEqual(host_state['players'], 2)
        self.assertEqual([p['nickname'] for p in host_state['joined']], ['Аня Ким', 'pupil'])
        self.assertIn(player_cookie(pin), guest.cookies)
        self.assertTrue(self.client.get(reverse('quiz_pin_lookup') + '?pin=' + pin).json()['ok'])
        self.assertFalse(self.client.get(reverse('quiz_pin_lookup') + '?pin=000000').json()['ok'])

    def test_bad_nicknames(self):
        pin = self.launch()
        self.join(pin, 'Айбек')
        self.join(pin, 'айбек', status=400)     # taken, any case
        self.join(pin, 'сука', status=400)
        self.join(pin, 'cyka', status=400)      # latin lookalikes
        self.join(pin, 'x' * 30)                # cut to 24 characters
        self.assertEqual(sorted(p.nickname for p in QuizPlayer.objects.all()), ['x' * 24, 'Айбек'])
        self.assertEqual(engine.player_count(self.game(pin).id), 2)
        response = Client().post(reverse('quiz_join'), {'pin': pin})
        self.assertEqual(response.status_code, 200)    # asks for a nickname
        self.assertEqual(Client().post(reverse('quiz_join'), {'pin': '999999', 'nickname': 'A'}).status_code, 404)

    def test_guests_can_be_turned_off(self):
        pin = self.launch(guests='off')
        self.join(pin, 'Гость', status=400)
        self.assertFalse(QuizPlayer.objects.exists())
        self.login(self.student).post(reverse('quiz_join'), {'pin': pin})
        self.assertEqual(QuizPlayer.objects.get().nickname, 'pupil')

    def test_reconnect_keeps_player_and_score(self):
        pin, (a, b) = self.game_with_players('Аня', 'Бек')
        user = self.login(self.student)
        user.post(reverse('quiz_join'), {'pin': pin})
        self.action(pin, 'start')
        right = option_id(self.q1, 'Два')
        for client in (a, b, user):
            self.assertEqual(self.answer(client, pin, 0, right).status_code, 200)
        self.assertEqual(self.phase(pin), 'reveal')
        anya = self.player(pin, 'Аня')
        self.assertEqual(anya.score, 1000)
        # same browser comes back
        response = a.get(reverse('quiz_join') + '?pin=' + pin)
        self.assertRedirects(response, reverse('quiz_play', args=[pin]), fetch_redirect_response=False)
        # player cookie lost, guest cookie kept
        again = Client()
        again.cookies['qz_guest'] = a.cookies['qz_guest'].value
        self.join(pin, 'Другой ник', client=again)
        self.assertEqual(self.result(again, pin).json()['me']['score'], 1000)
        # the same user on another device
        other_device = self.login(self.student)
        other_device.post(reverse('quiz_join'), {'pin': pin})
        self.assertEqual(self.result(other_device, pin).json()['me']['score'], 1000)
        self.assertEqual(QuizPlayer.objects.filter(game=self.game(pin)).count(), 3)
        self.assertEqual(engine.player_count(self.game(pin).id), 3)
        self.assertEqual(self.player(pin, 'Аня').id, anya.id)

    def test_kick(self):
        pin, (a, b) = self.game_with_players('Аня', 'Бек')
        bek = self.player(pin, 'Бек')
        self.action(pin, 'kick', player=bek.id)
        self.assertEqual(self.state(b, pin), {'kicked': True})
        self.assertNotIn('kicked', self.state(a, pin))
        self.assertEqual(engine.player_count(self.game(pin).id), 1)
        self.assertTrue(QuizPlayer.objects.get(id=bek.id).is_removed)
        self.join(pin, 'Бек снова', client=b, status=400)          # rejoin refused
        self.assertEqual(engine.player_count(self.game(pin).id), 1)
        self.action(pin, 'start')
        self.assertEqual(self.answer(b, pin, 0, option_id(self.q1, 'Два')).status_code, 403)
        self.assertEqual(self.answer(a, pin, 0, option_id(self.q1, 'Два')).status_code, 200)
        self.assertEqual(self.phase(pin), 'reveal')                 # everyone left has answered
        board = self.result(a, pin).json()['leaderboard']
        self.assertEqual(board['total'], 1)
        self.assertEqual([row['nick'] for row in board['top']], ['Аня'])
        self.action(pin, 'kick', player=bek.id)                     # twice: no double decrement
        self.assertEqual(engine.player_count(self.game(pin).id), 1)

    def test_kicked_answer_does_not_close_the_question_early(self):
        pin, (a, b, c) = self.game_with_players('Аня', 'Бек', 'Чынара')
        self.action(pin, 'start')
        self.assertEqual(self.answer(a, pin, 0, option_id(self.q1, 'Два')).status_code, 200)
        self.action(pin, 'kick', player=self.player(pin, 'Аня').id)
        self.assertEqual(self.answer(b, pin, 0, option_id(self.q1, 'Два')).status_code, 200)
        self.assertEqual(self.phase(pin), 'question')               # Чынара has not answered yet

    def test_player_limit(self):
        self.assertEqual(GameSession.MAX_PLAYERS, 1000)
        game = engine.create_game(self.quiz, self.author.profile, GameSession.HOST_SCREEN, {'max_players': 5000})
        self.assertEqual(game.settings['max_players'], 1000)
        game = engine.create_game(self.quiz, self.author.profile, GameSession.HOST_SCREEN)
        self.assertEqual(game.settings['max_players'], 1000)
        pin = self.launch()
        game = self.game(pin)
        game.settings = dict(game.settings, max_players=2)
        game.save()
        self.join(pin, 'Аня')
        self.join(pin, 'Бек')
        self.join(pin, 'Чынара', status=400)
        self.assertEqual(engine.player_count(game.id), 2)
        self.assertEqual(QuizPlayer.objects.filter(game=game).count(), 2)
        self.action(pin, 'kick', player=self.player(pin, 'Бек').id)
        self.join(pin, 'Чынара')


class QuizPhasesTest(QuizGameTestCase):
    def test_hosted_game_from_lobby_to_finish(self):
        pin = self.launch()
        state = self.state(Client(), pin)
        self.assertEqual((state['ph'], state['i'], state['n'], state['mode']), ('lobby', -1, 2, 'screen'))
        self.assertEqual(self.action(pin, 'start', status=409)['error'], 'Пока нет ни одного игрока')
        a, b = self.join(pin, 'Аня'), self.join(pin, 'Бек')
        self.assertEqual(self.result(a, pin).status_code, 409)
        state = self.action(pin, 'start')
        self.assertEqual((state['ph'], state['i']), ('question', 0))
        self.assertEqual(state['end'] - state['start'], self.q1.time_limit)
        self.assertEqual(self.action(pin, 'start')['ph'], 'question')    # second start does nothing
        right = option_id(self.q1, 'Два')
        self.answer(a, pin, 0, right)
        self.assertEqual(self.phase(pin), 'question')
        self.answer(b, pin, 0, option_id(self.q1, 'Три'))
        self.assertEqual(self.phase(pin), 'reveal')                       # everyone answered
        reveal = self.result(a, pin).json()['reveal']
        self.assertEqual(reveal['correct'], [right])
        self.assertEqual((reveal['answered'], reveal['players'], reveal['correct_count']), (2, 2, 1))
        self.assertEqual(reveal['counts'][str(right)], 1)
        self.assertEqual(self.action(pin, 'next')['ph'], 'leaderboard')
        self.assertEqual(self.result(a, pin).status_code, 200)
        state = self.action(pin, 'next')
        self.assertEqual((state['ph'], state['i']), ('question', 1))
        self.answer(a, pin, 1, option_id(self.q2, 'Да'))
        self.t += self.q2.time_limit                                       # not yet past the grace
        self.assertEqual(self.phase(pin), 'question')
        self.t += engine.GRACE + 0.1
        self.assertEqual(self.phase(pin), 'reveal')                        # closed by the poll
        self.assertEqual(self.action(pin, 'next')['ph'], 'leaderboard')
        self.assertEqual(self.action(pin, 'next')['ph'], 'finished')
        self.assertEqual(self.action(pin, 'next')['ph'], 'finished')
        game = self.game(pin)
        self.assertEqual(game.status, GameSession.FINISHED)
        self.assertIsNotNone(game.finished)
        self.quiz.refresh_from_db()
        self.assertEqual(self.quiz.plays, 1)
        self.join(pin, 'Опоздал', status=400)

    def test_auto_mode_runs_by_timers(self):
        pin, (a,) = self.game_with_players('Аня', mode='auto')
        self.action(pin, 'start')
        self.t += self.q1.time_limit + engine.GRACE + 0.1
        self.assertEqual(self.phase(pin), 'reveal')
        self.t += engine.AUTO_REVEAL - 1
        self.assertEqual(self.phase(pin), 'reveal')
        self.t += 1.1
        self.assertEqual(self.phase(pin), 'leaderboard')
        self.t += engine.AUTO_BOARD + 0.1
        state = self.state(a, pin)
        self.assertEqual((state['ph'], state['i']), ('question', 1))
        self.answer(a, pin, 1, option_id(self.q2, 'Да'))                 # the only player: closes at once
        self.assertEqual(self.phase(pin), 'reveal')
        self.t += engine.AUTO_REVEAL + 0.1
        self.assertEqual(self.phase(pin), 'leaderboard')
        self.t += engine.AUTO_BOARD + 0.1
        self.assertEqual(self.phase(pin), 'finished')

    def test_screen_mode_does_not_move_by_itself(self):
        pin, (a,) = self.game_with_players('Аня')
        self.action(pin, 'start')
        self.answer(a, pin, 0, option_id(self.q1, 'Два'))
        self.t += 600
        self.assertEqual(self.phase(pin), 'reveal')

    def test_pause_and_resume_shift_the_deadline(self):
        pin, (a, b) = self.game_with_players('Аня', 'Бек')
        state = self.action(pin, 'start')
        start, end = state['start'], state['end']
        self.t += 5
        self.assertEqual(self.action(pin, 'pause')['paused'], {'at': self.t})
        self.assertEqual(self.answer(a, pin, 0, option_id(self.q1, 'Два')).status_code, 409)
        self.t += 100                                    # far past the deadline, but paused
        self.assertEqual(self.phase(pin), 'question')
        state = self.action(pin, 'resume')
        self.assertIsNone(state['paused'])
        self.assertAlmostEqual(state['start'], start + 100)
        self.assertAlmostEqual(state['end'], end + 100)
        self.assertEqual(self.phase(pin), 'question')
        self.assertEqual(self.answer(a, pin, 0, option_id(self.q1, 'Два')).status_code, 200)
        self.assertEqual(QuizAnswer.objects.get().time_ms, 5000)

    def test_skip_and_end(self):
        pin, (a, b) = self.game_with_players('Аня', 'Бек')
        self.action(pin, 'start')
        self.answer(a, pin, 0, option_id(self.q1, 'Два'))
        self.assertEqual(self.action(pin, 'skip')['ph'], 'reveal')
        self.assertEqual(self.player(pin, 'Аня').score, 1000)
        self.assertEqual(self.action(pin, 'skip')['ph'], 'reveal')      # nothing to skip
        self.action(pin, 'next')
        self.action(pin, 'next')
        self.answer(b, pin, 1, option_id(self.q2, 'Да'))
        self.assertEqual(self.action(pin, 'end')['ph'], 'finished')      # the open question still counts
        self.assertEqual(self.player(pin, 'Бек').score, 1000)
        self.assertEqual(self.action(pin, 'end')['ph'], 'finished')
        self.assertEqual(self.action(pin, 'dance', status=400)['error'], 'Неизвестное действие')

    def test_solo_game(self):
        player = Client()
        response = player.post(reverse('quiz_solo', args=[self.quiz.key]), {'nickname': 'Соло'})
        self.assertEqual(response.status_code, 302)
        game = GameSession.objects.get(mode=GameSession.SOLO)
        pin = game.pin
        self.assertIsNone(game.host)
        self.assertEqual(game.settings['max_players'], 1)
        self.assertEqual(self.state(player, pin)['ph'], 'lobby')
        next_url = reverse('quiz_solo_next', args=[pin])
        self.assertEqual(Client().post(next_url).status_code, 403)
        self.assertEqual(player.get(next_url).status_code, 405)
        self.assertEqual(player.post(next_url).json()['ph'], 'question')
        self.question(player, pin)
        self.t += 2
        self.assertEqual(self.answer(player, pin, 0, option_id(self.q1, 'Два')).status_code, 200)
        self.assertEqual(self.state(player, pin)['ph'], 'reveal')         # right after the answer
        me = self.result(player, pin).json()['me']
        self.assertEqual((me['points'], me['rank'], me['total']), (950, 1, 1))
        self.t += 600
        self.assertEqual(self.state(player, pin)['ph'], 'reveal')         # waits for the player
        state = player.post(next_url).json()
        self.assertEqual((state['ph'], state['i']), ('question', 1))       # no leaderboard in solo
        self.answer(player, pin, 1, option_id(self.q2, 'Нет'))
        self.assertEqual(player.post(next_url).json()['ph'], 'finished')
        self.assertEqual(Client().post(reverse('quiz_join'), {'pin': pin, 'nickname': 'X'}).status_code, 404)
        self.assertEqual(self.host.post(reverse('quiz_solo_next', args=[self.launch()])).status_code, 404)
        records = quiz_records(self.quiz)
        self.assertEqual(records, [{'place': 1, 'nickname': 'Соло', 'score': 950, 'is_user': False}])
        user = self.login(self.student)
        user.post(reverse('quiz_solo', args=[self.quiz.key]))
        self.assertEqual(QuizPlayer.objects.get(profile=self.student.profile).nickname, 'pupil')
        self.assertEqual(self.client.get(reverse('quiz_solo', args=[self.quiz.key])).status_code, 302)

    @needs_templates('quiz/detail.html', 'quiz/records.html')
    def test_solo_records_on_pages(self):
        player = Client()
        player.post(reverse('quiz_solo', args=[self.quiz.key]), {'nickname': 'Рекордсмен'})
        pin = GameSession.objects.get(mode=GameSession.SOLO).pin
        next_url = reverse('quiz_solo_next', args=[pin])
        player.post(next_url)
        self.answer(player, pin, 0, option_id(self.q1, 'Два'))
        player.post(next_url)
        self.answer(player, pin, 1, option_id(self.q2, 'Да'))
        player.post(next_url)
        self.assertContains(self.client.get(reverse('quiz_detail', args=[self.quiz.key])), 'Рекордсмен')
        self.assertContains(self.client.get(reverse('quiz_records', args=[self.quiz.key])), 'Рекордсмен')


class QuizStateTest(QuizGameTestCase):
    def test_unchanged_version_gives_empty_answer(self):
        pin, (a,) = self.game_with_players('Аня')
        state = self.state(a, pin)
        self.assertEqual(self.state(a, pin, v=state['v']), {})
        self.assertEqual(self.state(Client(), pin, v=state['v']), {})
        self.action(pin, 'start')
        fresh = self.state(a, pin, v=state['v'])
        self.assertGreater(fresh['v'], state['v'])
        self.assertEqual(fresh['ph'], 'question')
        self.assertEqual(fresh['now'], round(self.t, 3))
        response = Client().get(reverse('quiz_state', args=['123456']))
        self.assertEqual(response.status_code, 404)
        self.assertTrue(response.json()['gone'])

    def test_state_polls_without_database(self):
        pin, (a, b) = self.game_with_players('Аня', 'Бек')
        self.action(pin, 'start')
        v = self.state(a, pin)['v']
        self.question(a, pin)
        anonymous = Client()
        with self.assertNumQueries(0):
            self.state(a, pin, v=v)
            self.state(a, pin)
            self.state(anonymous, pin)
            self.question(b, pin)
        with self.assertNumQueries(0):
            self.t += 3
            self.state(b, pin, v=v)

    def test_state_poll_is_one_cache_round_trip(self):
        pin, (a, b) = self.game_with_players('Аня', 'Бек')
        self.action(pin, 'start')
        v = self.state(a, pin)['v']
        calls = []
        real_get, real_get_many = cache.get, cache.get_many
        with patch.object(cache, 'get', side_effect=lambda *x, **y: calls.append('get') or real_get(*x, **y)), \
                patch.object(cache, 'get_many',
                             side_effect=lambda *x, **y: calls.append('get_many') or real_get_many(*x, **y)):
            self.assertEqual(self.state(a, pin, v=v), {})
        # One get_many of three keys (the local-memory backend runs it as one get per key, Redis as MGET).
        self.assertEqual(calls, ['get_many', 'get', 'get', 'get'])

    def test_removed_player_is_seen_by_the_poll(self):
        pin, (a, b) = self.game_with_players('Аня', 'Бек')
        self.action(pin, 'kick', player=self.player(pin, 'Аня').id)
        self.assertEqual(self.state(a, pin), {'kicked': True})
        self.assertNotIn('kicked', self.state(b, pin))

    def test_poll_interval(self):
        self.assertEqual(engine.poll_interval(100), 1000)
        self.assertEqual(engine.poll_interval(300), 1500)
        self.assertEqual(engine.poll_interval(800), 2000)
        self.assertEqual(engine.poll_interval(1000), 2000)
        pin, (a,) = self.game_with_players('Аня')
        game_id = self.game(pin).id
        for players, interval in ((100, 1000), (300, 1500), (800, 2000)):
            cache.set(engine.k('np', game_id), players)
            self.assertEqual(self.state(a, pin)['poll'], interval)


class QuizResultsTest(QuizGameTestCase):
    def test_leaderboard_places_and_streak(self):
        pin, (a, b, c) = self.game_with_players('Аня', 'Бек', 'Чынара')
        self.action(pin, 'start')
        for client in (a, b, c):
            self.question(client, pin)
        self.answer(a, pin, 0, option_id(self.q1, 'Два'))
        self.t += 5
        self.answer(b, pin, 0, option_id(self.q1, 'Два'))
        self.answer(c, pin, 0, option_id(self.q1, 'Три'))
        data = self.result(a, pin).json()
        self.assertEqual([(r['nick'], r['score'], r['rank']) for r in data['leaderboard']['top']],
                         [('Аня', 1000, 1), ('Бек', 875, 2), ('Чынара', 0, 3)])
        self.assertEqual(data['leaderboard']['total'], 3)
        me = data['me']
        self.assertEqual((me['answered'], me['correct'], me['points'], me['score'], me['rank'], me['total'],
                          me['streak']), (True, True, 1000, 1000, 1, 3, 1))
        me = self.result(c, pin).json()['me']
        self.assertEqual((me['correct'], me['points'], me['rank'], me['streak']), (False, 0, 3, 0))
        self.assertIsNone(self.result(Client(), pin).json()['me'])
        self.action(pin, 'next')
        self.action(pin, 'next')
        self.answer(a, pin, 1, option_id(self.q2, 'Да'))
        self.answer(b, pin, 1, option_id(self.q2, 'Нет'))
        self.answer(c, pin, 1, option_id(self.q2, 'Да'))
        me = self.result(a, pin).json()['me']
        self.assertEqual((me['points'], me['score'], me['streak'], me['rank']), (1100, 2100, 2, 1))
        top = self.result(c, pin).json()['leaderboard']['top']
        self.assertEqual([(r['nick'], r['rank'], r['prev'], r['delta']) for r in top],
                         [('Аня', 1, 1, 1100), ('Чынара', 2, 3, 1000), ('Бек', 3, 2, 0)])
        self.assertEqual(self.client.get(reverse('quiz_final', args=[pin])).status_code, 409)
        self.action(pin, 'next')
        self.action(pin, 'next')
        final = self.result(a, pin)
        self.assertEqual(final.status_code, 200)
        final = c.get(reverse('quiz_final', args=[pin])).json()
        self.assertEqual([(r['nickname'], r['score'], r['rank']) for r in final['top']],
                         [('Аня', 2100, 1), ('Чынара', 1000, 2), ('Бек', 875, 3)])
        self.assertEqual((final['total'], final['questions']), (3, 2))
        self.assertEqual(final['me'], {'nickname': 'Чынара', 'score': 1000, 'correct': 1, 'rank': 2})
        self.assertEqual(final['mistakes'], [{'n': 1, 'text': 'Вопрос single', 'type': 'single', 'given': 'Три',
                                              'correct': 'Два', 'explanation': ''}])
        self.assertEqual(a.get(reverse('quiz_final', args=[pin])).json()['mistakes'], [])
        self.assertIsNone(Client().get(reverse('quiz_final', args=[pin])).json()['me'])

    def test_mistakes_list_skips_unscored_and_shows_missing_answers(self):
        quiz = make_quiz(self.author)
        add_question(quiz, Q.TEXT, data={'answers': ['Бишкек', 'г. Бишкек']}, explanation='Столица')
        add_question(quiz, Q.POLL, [('x', False), ('y', False)])
        add_question(quiz, Q.ORDER, [('один', False), ('два', False), ('три', False)])
        add_question(quiz, Q.SINGLE, [('a', True), ('b', False)], weight=0)
        pin, (a,) = self.game_with_players('Аня', quiz=quiz)
        self.action(pin, 'start')
        self.answer(a, pin, 0, 'Ош')
        for _step in range(9):
            self.action(pin, 'skip')
            self.action(pin, 'next')
        self.assertEqual(self.phase(pin), 'finished')
        mistakes = a.get(reverse('quiz_final', args=[pin])).json()['mistakes']
        self.assertEqual([(m['n'], m['given'], m['correct']) for m in mistakes],
                         [(1, 'Ош', 'Бишкек / г. Бишкек'), (3, None, 'один → два → три')])
        self.assertEqual(mistakes[0]['explanation'], 'Столица')

    def test_double_and_zero_weight_and_partial(self):
        quiz = make_quiz(self.author)
        double = add_question(quiz, Q.SINGLE, [('a', True), ('b', False)], weight=2)
        zero = add_question(quiz, Q.SINGLE, [('a', True), ('b', False)], weight=0)
        multi = add_question(quiz, Q.MULTI, [('a', True), ('b', True), ('c', False)])
        pin, (a,) = self.game_with_players('Аня', quiz=quiz)
        self.action(pin, 'start')
        self.answer(a, pin, 0, option_id(double, 'a'))
        self.assertEqual(self.result(a, pin).json()['me']['points'], 2000)
        self.action(pin, 'next')
        self.action(pin, 'next')
        self.answer(a, pin, 1, option_id(zero, 'a'))
        me = self.result(a, pin).json()['me']
        self.assertEqual((me['points'], me['scored'], me['streak']), (0, False, 1))
        self.action(pin, 'next')
        self.action(pin, 'next')
        self.answer(a, pin, 2, [option_id(multi, 'a')])
        me = self.result(a, pin).json()['me']
        self.assertEqual((me['points'], me['fraction'], me['correct'], me['streak']), (500, 0.5, False, 0))


class QuizAccessTest(QuizGameTestCase):
    def test_host_endpoints_are_for_the_host(self):
        pin, (a,) = self.game_with_players('Аня')
        state_url = reverse('quiz_host_state', args=[pin])
        for client in (Client(), self.login(self.other), a):
            self.assertEqual(client.get(state_url).status_code, 403)
            self.assertEqual(client.post(reverse('quiz_host_action', args=[pin]), {'action': 'start'}).status_code,
                             403)
        self.assertEqual(self.phase(pin), 'lobby')
        self.assertEqual(self.login(self.other).get(reverse('quiz_host', args=[pin])).status_code, 403)
        self.assertEqual(Client().get(reverse('quiz_host', args=[pin])).status_code, 302)
        admin = self.login(self.admin)
        self.assertEqual(admin.get(state_url).status_code, 200)
        self.assertEqual(self.action(pin, 'start', client=admin)['ph'], 'question')
        data = self.host.get(state_url).json()
        self.assertEqual(data['key']['correct'], [option_id(self.q1, 'Два')])     # the host sees the key

    def test_launch_needs_login_and_post(self):
        url = reverse('quiz_launch', args=[self.quiz.key])
        response = Client().post(url, {'mode': 'screen'})
        self.assertEqual(response.status_code, 302)
        self.assertIn('login', response['Location'])
        self.assertEqual(self.host.get(url).status_code, 405)
        self.assertFalse(GameSession.objects.exists())
        self.launch(mode='solo')                       # not a hosted mode: falls back to the screen
        self.assertEqual(GameSession.objects.get().mode, GameSession.HOST_SCREEN)

    def test_hidden_quiz(self):
        self.quiz.is_hidden = True
        self.quiz.save()
        other = self.login(self.other)
        self.assertEqual(other.post(reverse('quiz_launch', args=[self.quiz.key])).status_code, 404)
        self.assertEqual(Client().get(reverse('quiz_detail', args=[self.quiz.key])).status_code, 404)
        self.assertEqual(Client().post(reverse('quiz_solo', args=[self.quiz.key])).status_code, 404)
        self.assertEqual(Client().get(reverse('quiz_records', args=[self.quiz.key])).status_code, 404)
        self.assertNotIn(self.quiz, visible_quizzes())
        self.launch()                                  # the author still can
        self.launch(client=self.login(self.admin))
        self.assertEqual(Client().get(reverse('quiz_detail', args=['nosuchkey'])).status_code, 404)

    def test_link_only_quiz(self):
        hidden_from_catalog = make_quiz(self.author, title='Только по ссылке', access=Quiz.LINK)
        add_question(hidden_from_catalog, Q.SINGLE, [('a', True), ('b', False)])
        empty = make_quiz(self.author, title='Пустая')
        listed = list(visible_quizzes())
        self.assertIn(self.quiz, listed)
        self.assertNotIn(hidden_from_catalog, listed)
        self.assertNotIn(empty, listed)
        self.launch(quiz=hidden_from_catalog, client=self.login(self.other))   # opens by key

    @needs_templates('quiz/catalog.html', 'quiz/detail.html')
    def test_link_only_quiz_pages(self):
        hidden_from_catalog = make_quiz(self.author, title='Только по ссылке', access=Quiz.LINK)
        add_question(hidden_from_catalog, Q.SINGLE, [('a', True), ('b', False)])
        catalog = self.client.get(reverse('quiz_catalog'))
        self.assertContains(catalog, 'Столицы')
        self.assertNotContains(catalog, 'Только по ссылке')
        self.assertContains(self.client.get(reverse('quiz_detail', args=[hidden_from_catalog.key])),
                            'Только по ссылке')

    @needs_templates('quiz/detail.html')
    def test_hidden_quiz_page_for_author(self):
        self.quiz.is_hidden = True
        self.quiz.save()
        self.assertEqual(self.host.get(reverse('quiz_detail', args=[self.quiz.key])).status_code, 200)
        self.assertEqual(self.login(self.other).get(reverse('quiz_detail', args=[self.quiz.key])).status_code, 404)

    def played_game(self):
        pin, (a, b) = self.game_with_players('Аня', 'Бек')
        self.action(pin, 'start')
        self.answer(a, pin, 0, option_id(self.q1, 'Два'))
        self.answer(b, pin, 0, option_id(self.q1, 'Три'))
        self.action(pin, 'end')
        return self.game(pin)

    def test_xlsx_report(self):
        from openpyxl import load_workbook

        game = self.played_game()
        url = reverse('quiz_game_report_xlsx', args=[game.id])
        self.assertEqual(self.login(self.other).get(url).status_code, 404)
        self.assertEqual(self.login(self.student).get(reverse('quiz_game_report', args=[game.id])).status_code,
                         404)
        self.assertEqual(Client().get(url).status_code, 302)
        response = self.host.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'],
                         'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book.sheetnames, ['Игроки', 'Вопросы'])
        rows = list(book['Игроки'].iter_rows(values_only=True))
        self.assertEqual(rows[0][:7], ('Место', 'Ник', 'Пользователь', 'Очки', 'Верных', 'В1', 'В2'))
        self.assertEqual([rows[1][i] for i in (0, 1, 3, 4, 5)], [1, 'Аня', 1000, 1, 1000])
        self.assertEqual([rows[2][i] for i in (0, 1, 3, 4, 5)], [2, 'Бек', 0, 0, 0])
        self.assertFalse(rows[1][2])                     # guests have no username
        questions = list(book['Вопросы'].iter_rows(values_only=True))
        self.assertEqual(questions[1][3:6], (2, 1, 50))
        self.assertEqual(self.login(self.admin).get(url).status_code, 200)

    def test_report_for_the_quiz_author(self):
        game = engine.create_game(self.quiz, self.other.profile, GameSession.HOST_SCREEN)
        url = reverse('quiz_game_report_xlsx', args=[game.id])
        self.assertEqual(self.host.get(url).status_code, 200)              # author of the quiz
        self.assertEqual(self.login(self.other).get(url).status_code, 200)  # host of the game
        self.assertEqual(self.login(self.student).get(url).status_code, 404)

    @needs_templates('quiz/report.html')
    def test_html_report(self):
        game = self.played_game()
        response = self.host.get(reverse('quiz_game_report', args=[game.id]))
        self.assertContains(response, 'Аня')
        self.assertEqual(self.login(self.other).get(reverse('quiz_game_report', args=[game.id])).status_code, 404)

    def test_games_can_be_disabled(self):
        MiscConfig.objects.create(key='quiz_disabled', value='Идёт олимпиада')
        self.assertEqual(engine.games_disabled(), 'Идёт олимпиада')
        self.launch(client=self.login(self.admin))
        self.assertEqual(GameSession.objects.count(), 1)

    @needs_templates('quiz/message.html')
    def test_disabled_games_stop_launch(self):
        MiscConfig.objects.create(key='quiz_disabled', value='Идёт олимпиада')
        response = self.host.post(reverse('quiz_launch', args=[self.quiz.key]), {'mode': 'screen'})
        self.assertEqual(response.status_code, 403)
        self.assertContains(response, 'Идёт олимпиада', status_code=403)
        self.assertEqual(Client().post(reverse('quiz_solo', args=[self.quiz.key]), {'nickname': 'Соло'})
                         .status_code, 403)
        self.assertFalse(GameSession.objects.exists())

    @needs_templates('quiz/catalog.html', 'quiz/detail.html', 'quiz/my.html')
    def test_pages_render(self):
        self.assertContains(self.client.get(reverse('quiz_catalog')), 'Столицы')
        self.assertEqual(self.client.get(reverse('quiz_catalog') + '?q=Стол&order=popular').status_code, 200)
        self.assertContains(self.client.get(reverse('quiz_detail', args=[self.quiz.key])), 'Столицы')
        self.assertEqual(self.client.get(reverse('quiz_my')).status_code, 302)
        self.played_game()
        self.assertContains(self.host.get(reverse('quiz_my')), 'Столицы')
        self.assertEqual(self.login(self.student).get(reverse('quiz_my')).status_code, 200)

    @needs_templates('quiz/host.html')
    def test_one_hosted_mode_with_pin_and_qr_in_both_layouts(self):
        self.assertEqual([m for m, _name in GameSession.MODES], ['screen', 'auto', 'solo'])
        pin = self.launch(mode='remote')                 # the old mode name still starts a hosted game
        self.assertEqual(self.game(pin).mode, GameSession.HOST)
        qr = reverse('quiz_qr', args=[pin])
        for view in ('', '?view=screen', '?view=remote'):
            page = self.host.get(reverse('quiz_host', args=[pin]) + view).content.decode()
            self.assertIn(pin, page)
            self.assertIn(qr, page)

    @needs_templates('quiz/play.html', 'quiz/host.html')
    def test_game_pages_render(self):
        pin, (a,) = self.game_with_players('Аня')
        self.assertEqual(a.get(reverse('quiz_play', args=[pin])).status_code, 200)
        self.assertEqual(self.host.get(reverse('quiz_host', args=[pin])).status_code, 200)
        self.assertEqual(self.host.get(reverse('quiz_host', args=[pin]) + '?view=remote').status_code, 200)
        self.assertEqual(Client().get(reverse('quiz_play', args=[pin])).status_code, 302)

    def test_qr_code(self):
        pin = self.launch()
        response = self.client.get(reverse('quiz_qr', args=[pin]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'image/svg+xml')
        self.assertIn(b'<svg', response.content)


# ---------------------------------------------------------------- the editor

def png_bytes(size=(2000, 1500), mode='RGB'):
    from PIL import Image

    image = Image.new(mode, size, (30, 120, 200) if mode == 'RGB' else (30, 120, 200, 128))
    buf = BytesIO()
    image.save(buf, 'PNG')
    return buf.getvalue()


class QuizEditorTest(TestCase):
    def setUp(self):
        cache.clear()
        self.media = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.media, ignore_errors=True)
        media = override_settings(MEDIA_ROOT=self.media)
        media.enable()
        self.addCleanup(media.disable)
        self.author = create_user(username='editor')
        self.other = create_user(username='neighbour')
        self.editor = Client()
        self.editor.force_login(self.author)
        self.stranger = Client()
        self.stranger.force_login(self.other)
        self.quiz = make_quiz(self.author, title='Моя', access=Quiz.LINK)

    def api(self, name, *args, client=None, quiz=None, body=None, files=None, method='post'):
        url = reverse(name, args=[(quiz or self.quiz).key] + list(args))
        client = client or self.editor
        if method == 'get':
            return client.get(url)
        if files is not None:
            return client.post(url, files)
        return client.post(url, json.dumps(body or {}), content_type='application/json')

    def add(self, type, **kwargs):
        response = self.api('quiz_api_question_add', body={'type': type}, **kwargs)
        self.assertEqual(response.status_code, 200, response.content[:300])
        return response.json()['question']

    def save(self, question_id, payload, status=200, **kwargs):
        response = self.api('quiz_api_question_save', question_id, body=payload, **kwargs)
        self.assertEqual(response.status_code, status, response.content[:300])
        return response.json()

    def test_new_quiz(self):
        response = self.editor.get(reverse('quiz_new'))
        quiz = Quiz.objects.exclude(id=self.quiz.id).get()
        self.assertRedirects(response, reverse('quiz_edit', args=[quiz.key]), fetch_redirect_response=False)
        self.assertEqual(quiz.author, self.author.profile)
        self.assertRegex(quiz.key, r'^[a-z0-9]{8}$')
        self.assertEqual(self.editor.get(reverse('quiz_new')).status_code, 302)    # the empty one is reused
        self.assertEqual(Quiz.objects.count(), 2)
        self.assertEqual(Client().get(reverse('quiz_new')).status_code, 302)
        self.assertEqual(Quiz.objects.count(), 2)

    def test_only_the_author_edits(self):
        question = self.add(Q.SINGLE)
        payload = {'type': Q.SINGLE, 'text': 'Взлом', 'options': [{'text': 'a', 'is_correct': True}]}
        self.assertEqual(self.save(question['id'], payload, status=404, client=self.stranger)['error'],
                         'Нет доступа к этой викторине')
        self.assertEqual(self.api('quiz_api_question_add', client=self.stranger, body={'type': Q.SINGLE})
                         .status_code, 404)
        self.assertEqual(self.api('quiz_api_save', client=self.stranger, body={'title': 'Чужая'}).status_code, 404)
        self.assertEqual(self.api('quiz_api_question_delete', question['id'], client=self.stranger).status_code,
                         404)
        self.assertEqual(self.api('quiz_api_data', client=self.stranger, method='get').status_code, 404)
        self.assertEqual(self.api('quiz_api_question_save', question['id'], client=Client(),
                                  body=payload).status_code, 403)
        self.assertEqual(self.stranger.get(reverse('quiz_edit', args=[self.quiz.key])).status_code, 403)
        self.assertEqual(self.stranger.post(reverse('quiz_delete', args=[self.quiz.key])).status_code, 403)
        self.quiz.is_hidden = True
        self.quiz.save()
        self.assertEqual(self.stranger.get(reverse('quiz_edit', args=[self.quiz.key])).status_code, 404)
        self.assertEqual(Client().get(reverse('quiz_edit', args=[self.quiz.key])).status_code, 302)
        self.quiz.refresh_from_db()
        self.assertEqual(self.quiz.title, 'Моя')
        self.assertEqual(QuizQuestion.objects.get().text, '')
        # a question of another quiz cannot be reached through one's own quiz
        mine = make_quiz(self.other)
        self.assertEqual(self.save(question['id'], payload, status=404, client=self.stranger, quiz=mine)['error'],
                         'Вопрос не найден')

    def test_author_and_admin_edit(self):
        admin = Client()
        admin.force_login(create_user(username='editadmin', is_superuser=True, is_staff=True))
        response = self.api('quiz_api_save', client=admin, body={'title': 'Чужое', 'access': 'public'})
        self.assertEqual(response.status_code, 200)
        response = self.api('quiz_api_save', body={'title': '  Новое   имя ', 'access': 'public'})
        self.assertEqual(response.status_code, 200)
        self.quiz.refresh_from_db()
        self.assertEqual((self.quiz.title, self.quiz.access), ('Новое имя', Quiz.PUBLIC))
        self.assertEqual(self.api('quiz_api_save', body={'title': ''}).json()['field'], 'title')
        self.assertEqual(self.api('quiz_api_save', body={'theme': 'neon'}).json()['field'], 'theme')
        self.assertEqual(self.api('quiz_api_save', body={'access': 'all'}).json()['field'], 'access')
        self.editor.post(reverse('quiz_delete', args=[self.quiz.key]))
        self.assertFalse(Quiz.objects.filter(id=self.quiz.id).exists())

    def test_every_type_is_saved_as_the_game_expects(self):
        cases = [
            (Q.SINGLE, {'options': [{'text': 'a', 'is_correct': True}, {'text': 'b', 'is_correct': True},
                                    {'text': 'c'}]}, {}),
            (Q.TRUEFALSE, {'options': [{'text': 'x', 'is_correct': False}, {'text': 'y', 'is_correct': True}]}, {}),
            (Q.MULTI, {'options': [{'text': 'a', 'is_correct': True}, {'text': 'b', 'is_correct': True},
                                   {'text': 'c'}]}, {}),
            (Q.WORD, {'data': {'word': ' алгоритм '}}, {'word': 'АЛГОРИТМ'}),
            (Q.ORDER, {'options': [{'text': '1', 'is_correct': True}, {'text': '2'}, {'text': '3'}]}, {}),
            (Q.CODE, {'code': 'print(5)', 'code_language': 'python',
                      'options': [{'text': '5', 'is_correct': True}, {'text': '6'}], 'data': {'input': False}},
             {'input': False}),
            (Q.CODE, {'code': 'print(5)', 'data': {'input': True, 'answers': '5; 5.0', 'tolerance': '0'}},
             {'input': True, 'answers': ['5', '5.0'], 'tolerance': 0}),
            (Q.TEXT, {'data': {'answers': ['Бишкек', 'г. Бишкек', 'Бишкек'], 'tolerance': '0,5'}},
             {'answers': ['Бишкек', 'г. Бишкек'], 'tolerance': 0.5}),
            (Q.SLIDER, {'data': {'min': 0, 'max': '100', 'step': 1, 'answer': 42, 'tolerance': 2, 'unit': 'км'}},
             {'min': 0, 'max': 100, 'step': 1, 'answer': 42, 'tolerance': 2, 'unit': 'км'}),
            (Q.PIN, {'data': {'x': 30.5, 'y': 160, 'r': 8}}, {'x': 30.5, 'y': 100, 'r': 8}),
            (Q.POLL, {'options': [{'text': 'x', 'is_correct': True}, {'text': 'y'}]}, {}),
            (Q.CLOUD, {}, {}),
            (Q.OPEN, {}, {}),
        ]
        for type, extra, data in cases:
            with self.subTest(type=type, data=data):
                question = self.add(type)
                payload = dict({'type': type, 'text': 'Вопрос', 'time_limit': 30, 'weight': 2}, **extra)
                saved = self.save(question['id'], payload)['question']
                self.assertEqual(saved['data'], data)
                stored = QuizQuestion.objects.get(id=question['id'])
                self.assertEqual((stored.type, stored.time_limit, stored.weight), (type, 30, 2))
                correct = [o.is_correct for o in stored.options.all()]
                if type == Q.SINGLE:
                    self.assertEqual(correct, [True, False, False])        # only one right option
                elif type == Q.TRUEFALSE:
                    self.assertEqual([o.text for o in stored.options.all()], ['Верно', 'Неверно'])
                    self.assertEqual(correct, [False, True])
                elif type in (Q.ORDER, Q.POLL):
                    self.assertFalse(any(correct))
                elif type in (Q.WORD, Q.TEXT, Q.SLIDER, Q.PIN, Q.CLOUD, Q.OPEN) or data.get('input'):
                    self.assertEqual(correct, [])
                scoring.build_key(stored)                                  # the game can check it
        self.assertEqual(self.quiz.questions.count(), len(cases))
        self.assertEqual(list(self.quiz.questions.values_list('order', flat=True)), list(range(len(cases))))
        pin = self.api('quiz_api_data', method='get').json()['questions'][9]
        self.assertEqual(pin['type'], Q.PIN)
        self.assertIn('Загрузите картинку', pin['problems'])

    def test_validation_errors(self):
        question = self.add(Q.SINGLE)
        self.assertIn('Отметьте верный вариант', question['problems'])
        self.assertEqual(self.api('quiz_api_question_add', body={'type': 'riddle'}).json()['field'], 'type')
        base = {'type': Q.SINGLE, 'text': 'Вопрос', 'options': [{'text': 'a', 'is_correct': True}, {'text': 'b'}]}
        for change, field in (({'time_limit': 3}, 'time_limit'), ({'time_limit': 500}, 'time_limit'),
                              ({'time_limit': 'долго'}, 'time_limit'), ({'weight': 5}, 'weight'),
                              ({'type': 'riddle'}, 'type'), ({'text': 'я' * 2001}, 'text'),
                              ({'options': [{'text': str(n)} for n in range(5)]}, 'options'),
                              ({'type': Q.WORD, 'data': {'word': 'а' * 30}}, 'word'),
                              ({'type': Q.SLIDER, 'data': {'min': 10, 'max': 5}}, 'max'),
                              ({'type': Q.SLIDER, 'data': {'min': 0, 'max': 10, 'answer': 50}}, 'answer'),
                              ({'type': Q.SLIDER, 'data': {'min': 0, 'max': 10, 'step': 0}}, 'step'),
                              ({'type': Q.TEXT, 'data': {'answers': ['1'], 'tolerance': 'nan'}}, 'tolerance')):
            with self.subTest(change=str(change)[:60]):
                self.assertEqual(self.save(question['id'], dict(base, **change), status=400)['field'], field)
        self.assertEqual(QuizQuestion.objects.get().type, Q.SINGLE)
        url = reverse('quiz_api_question_save', args=[self.quiz.key, question['id']])
        self.assertEqual(self.editor.post(url, 'not json', content_type='application/json').status_code, 400)
        self.assertEqual(self.editor.post(url, '[1]', content_type='application/json').status_code, 400)
        self.assertEqual(self.editor.get(url).status_code, 405)

    def upload(self, question_id, content, name='picture.png', **data):
        upload = SimpleUploadedFile(name, content, content_type='image/png')
        return self.api('quiz_api_question_image', question_id, files=dict(data, image=upload))

    def test_image_upload(self):
        from PIL import Image

        question = self.add(Q.PIN)
        response = self.upload(question['id'], png_bytes())
        self.assertEqual(response.status_code, 200, response.content[:300])
        stored = QuizQuestion.objects.get(id=question['id'])
        self.assertTrue(stored.image.name.startswith('quiz/'))
        self.assertNotIn('picture', stored.image.name)
        with stored.image.open('rb') as f:
            content = f.read()
        self.assertLessEqual(len(content), 100 * 1024)
        self.assertEqual(Image.open(BytesIO(content)).size, (1280, 960))
        self.assertNotIn('Загрузите картинку', response.json()['question']['problems'])
        # a transparent picture keeps its transparency
        option_question = self.add(Q.SINGLE)
        option = option_question['options'][0]['id']
        response = self.upload(option_question['id'], png_bytes((300, 200), 'RGBA'), option=option)
        self.assertEqual(response.status_code, 200)
        self.assertTrue(QuizOption.objects.get(id=option).image.name.endswith('.webp'))
        response = self.api('quiz_api_question_image', question['id'], files={'remove': '1'})
        self.assertEqual(response.status_code, 200)
        self.assertFalse(QuizQuestion.objects.get(id=question['id']).image)

    def test_bad_images_are_rejected(self):
        question = self.add(Q.PIN)
        response = self.upload(question['id'], b'just some text, not a picture')
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()['field'], 'image')
        self.assertEqual(self.upload(question['id'], png_bytes()[:200]).status_code, 400)        # cut short
        response = self.upload(question['id'], b'\x89PNG' + b'0' * (5 * 1024 * 1024 + 10))
        self.assertEqual(response.status_code, 400)
        self.assertIn('5 МБ', response.json()['error'])
        response = self.upload(question['id'], b'<svg xmlns="http://www.w3.org/2000/svg"></svg>', name='a.svg')
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.api('quiz_api_question_image', question['id'], files={}).status_code, 400)
        truefalse = self.add(Q.TRUEFALSE)
        response = self.upload(truefalse['id'], png_bytes((50, 50)), option=truefalse['options'][0]['id'])
        self.assertEqual(response.status_code, 400)
        self.assertFalse(QuizQuestion.objects.get(id=question['id']).image)
        self.assertEqual(os.listdir(self.media), [])

    def test_excel_import(self):
        from openpyxl import load_workbook

        template = self.client.get(reverse('quiz_import_template'))
        self.assertEqual(template.status_code, 200)
        content = template.content
        response = self.api('quiz_api_import', files={'file': SimpleUploadedFile('q.xlsx', content)})
        self.assertEqual(response.status_code, 200, response.content[:300])
        result = response.json()
        self.assertEqual(result['errors'], [])
        types = list(self.quiz.questions.values_list('type', flat=True))
        self.assertEqual(result['added'], len(types))
        self.assertEqual(types, [Q.SINGLE, Q.MULTI, Q.TRUEFALSE, Q.TEXT, Q.CODE, Q.WORD, Q.ORDER, Q.SLIDER,
                                 Q.POLL])
        single = self.quiz.questions.get(type=Q.SINGLE)
        self.assertEqual([o.text for o in single.options.filter(is_correct=True)], ['Бишкек'])
        multi = self.quiz.questions.get(type=Q.MULTI)
        self.assertEqual((multi.weight, [o.text for o in multi.options.filter(is_correct=True)]), (2, ['2', '11']))
        self.assertEqual(self.quiz.questions.get(type=Q.WORD).data, {'word': 'АЛГОРИТМ'})
        self.assertEqual(self.quiz.questions.get(type=Q.TEXT).data['answers'], ['56'])
        self.assertEqual(self.quiz.questions.get(type=Q.SLIDER).data['answer'], 7134)
        self.assertEqual(self.quiz.questions.get(type=Q.POLL).weight, 0)

        book = load_workbook(BytesIO(content))
        sheet = book['Вопросы']
        sheet.delete_rows(2, sheet.max_row)
        sheet.append(['Один вариант', 'Хороший', '', '', 'a', 'b', '', '', '', '', '1'])
        sheet.append(['Загадка', 'Что это?'])
        sheet.append(['Один вариант', 'Нет такого номера', '', '', 'a', 'b', '', '', '', '', '5'])
        sheet.append(['Точка на картинке', 'Где Ош?'])
        sheet.append(['Введи ответ', 'Долго', '', '', '', '', '', '', '', '', '1', 999])
        sheet.append([])
        sheet.append(['Ползунок', 'Без шкалы', '', '', '', '', '', '', '', '', 5, 20, '', '', 10, 1])
        buf = BytesIO()
        book.save(buf)
        response = self.api('quiz_api_import', files={'file': SimpleUploadedFile('q.xlsx', buf.getvalue())})
        result = response.json()
        self.assertEqual(result['added'], 1)
        self.assertEqual([e['row'] for e in result['errors']], [3, 4, 5, 6, 8])
        self.assertEqual(self.quiz.questions.count(), 10)
        response = self.api('quiz_api_import', files={'file': SimpleUploadedFile('q.xlsx', b'not excel')})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.api('quiz_api_import', files={}).status_code, 400)
        response = self.api('quiz_api_import', client=self.stranger,
                            files={'file': SimpleUploadedFile('q.xlsx', content)})
        self.assertEqual(response.status_code, 404)
        self.assertEqual(self.quiz.questions.count(), 10)

    def test_copy(self):
        question = add_question(self.quiz, Q.SINGLE, [('a', True), ('b', False)], explanation='почему')
        self.assertEqual(self.upload(question.id, png_bytes((100, 100))).status_code, 200)
        url = reverse('quiz_copy', args=[self.quiz.key])
        self.assertEqual(self.stranger.post(url).status_code, 404)                  # link-only: not copied
        self.quiz.access = Quiz.PUBLIC
        self.quiz.save()
        self.assertEqual(Client().post(url).status_code, 302)
        self.assertEqual(self.stranger.get(url).status_code, 405)
        response = self.stranger.post(url)
        copy = Quiz.objects.get(author=self.other.profile)
        self.assertRedirects(response, reverse('quiz_edit', args=[copy.key]), fetch_redirect_response=False)
        self.assertEqual((copy.copied_from, copy.access), (self.quiz, Quiz.PUBLIC))
        self.assertTrue(copy.title.endswith('Моя'))
        copied = copy.questions.get()
        self.assertEqual((copied.text, copied.explanation), (question.text, 'почему'))
        self.assertEqual([(o.text, o.is_correct) for o in copied.options.all()], [('a', True), ('b', False)])
        self.assertTrue(copied.image)
        self.assertNotEqual(copied.image.name, QuizQuestion.objects.get(id=question.id).image.name)
        self.assertTrue(copy.can_edit(self.other))
        self.assertEqual(self.quiz.questions.count(), 1)
        self.quiz.is_hidden = True
        self.quiz.save()
        self.assertEqual(self.stranger.post(url).status_code, 404)
        self.assertEqual(self.editor.post(url).status_code, 302)                    # own hidden quiz

    @needs_templates('quiz/editor.html')
    def test_editor_page(self):
        self.add(Q.SINGLE)
        self.assertEqual(self.editor.get(reverse('quiz_edit', args=[self.quiz.key])).status_code, 200)


class QuizRedisCacheTest(TestCase):
    def test_one_client_per_process(self):
        from judge.utils.redis_cache import RedisCache
        backend = RedisCache('redis://127.0.0.1:6379/9', {})
        client = backend._cache
        self.assertIs(client.get_client(), client.get_client())
        self.assertIs(client.get_client(write=True), client.get_client())
        try:
            backend.set('qz:test', {'a': 1}, 5)
        except Exception:
            self.skipTest('no Redis here')
        self.assertEqual(backend.get_many(['qz:test', 'qz:none']), {'qz:test': {'a': 1}})
        backend.delete('qz:test')

    def test_settings_swap_django_redis_for_ours(self):
        source = open(os.path.join(os.path.dirname(__file__), '..', 'dmoj', 'settings.py'), encoding='utf-8').read()
        start = source.index('for _cache in CACHES.values():')
        caches = {'default': {'BACKEND': 'django.core.cache.backends.redis.RedisCache', 'LOCATION': 'redis://x/1'},
                  'other': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache'}}
        exec(source[start:source.index('\n\n', start)], {'CACHES': caches})
        self.assertEqual(caches['default']['BACKEND'], 'judge.utils.redis_cache.RedisCache')
        self.assertEqual(caches['default']['LOCATION'], 'redis://x/1')
        self.assertEqual(caches['other']['BACKEND'], 'django.core.cache.backends.locmem.LocMemCache')
